"""
Tests for backend/app/services/extraction.py.

Covers:
- PDF text extraction with real patents from docs/
- DOCX extraction path
- TXT extraction path
- Unsupported file type guard
- OCR fallback threshold (MIN_CHARS_PER_PAGE)

Note: US11941918.pdf is a scanned (image-only) PDF whose text extraction
falls back to Tesseract OCR via pdf2image / Poppler. Tests that require it
are automatically skipped when Poppler is not found in PATH (e.g. on a fresh
Windows dev machine without Poppler installed).
"""

from pathlib import Path

import pytest

from app.services.extraction import MIN_CHARS_PER_PAGE, extract_text


def _poppler_available() -> bool:
    """Return True if poppler's pdfinfo binary is on PATH."""
    import subprocess
    try:
        subprocess.run(["pdfinfo", "-v"], capture_output=True, timeout=5)
        return True
    except (FileNotFoundError, OSError):
        return False


requires_poppler = pytest.mark.skipif(
    not _poppler_available(),
    reason="Poppler (pdfinfo) not installed / not in PATH — OCR-dependent tests skipped",
)

# Real patent PDFs that live in docs/ at the repo root
DOCS = Path(__file__).resolve().parents[2] / "docs"
PDF_US10156900 = DOCS / "US10156900.pdf"
PDF_US11941918 = DOCS / "US11941918.pdf"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _pdf_bytes(path: Path) -> bytes:
    if not path.exists():
        pytest.skip(f"PDF not found: {path}")
    return path.read_bytes()


# ---------------------------------------------------------------------------
# PDF extraction
# ---------------------------------------------------------------------------

class TestPdfExtraction:
    def test_us10156900_returns_nonempty_text(self):
        data = _pdf_bytes(PDF_US10156900)
        text = extract_text(data, "US10156900.pdf")
        assert isinstance(text, str)
        assert len(text) > 500, "Expected substantial text from US10156900"

    @requires_poppler
    def test_us11941918_returns_nonempty_text(self):
        data = _pdf_bytes(PDF_US11941918)
        text = extract_text(data, "US11941918.pdf")
        assert isinstance(text, str)
        assert len(text) > 500, "Expected substantial text from US11941918"

    def test_us10156900_contains_patent_keywords(self):
        """Real patents should mention standard patent-document sections."""
        data = _pdf_bytes(PDF_US10156900)
        text = extract_text(data, "US10156900.pdf").lower()
        # At least one of these common patent terms should appear
        found = any(kw in text for kw in ("claim", "abstract", "invention", "embodiment", "patent"))
        assert found, f"No patent keywords found in US10156900 text (first 200 chars): {text[:200]}"

    @requires_poppler
    def test_us11941918_contains_patent_keywords(self):
        data = _pdf_bytes(PDF_US11941918)
        text = extract_text(data, "US11941918.pdf").lower()
        found = any(kw in text for kw in ("claim", "abstract", "invention", "embodiment", "patent"))
        assert found, f"No patent keywords found in US11941918 text (first 200 chars): {text[:200]}"

    def test_pdf_extension_case_insensitive(self):
        """'.PDF' (uppercase) should be handled the same as '.pdf'."""
        data = _pdf_bytes(PDF_US10156900)
        text = extract_text(data, "US10156900.PDF")
        assert len(text) > 0

    @requires_poppler
    def test_both_pdfs_produce_distinct_text(self):
        """Two different patents should not produce identical extracted text."""
        t1 = extract_text(_pdf_bytes(PDF_US10156900), "US10156900.pdf")
        t2 = extract_text(_pdf_bytes(PDF_US11941918), "US11941918.pdf")
        assert t1 != t2, "Different PDFs should produce different text"


# ---------------------------------------------------------------------------
# TXT extraction
# ---------------------------------------------------------------------------

class TestTxtExtraction:
    def test_txt_roundtrip(self):
        content = "Hello, this is a test patent document.\nLine 2."
        result = extract_text(content.encode("utf-8"), "test.txt")
        assert result == content

    def test_txt_handles_latin1_gracefully(self):
        """Bytes that are valid latin-1 but not utf-8 should not raise."""
        data = b"Caf\xe9 patent"  # \xe9 is not valid utf-8
        result = extract_text(data, "doc.txt")
        assert isinstance(result, str)
        assert len(result) > 0

    def test_txt_extension_case_insensitive(self):
        data = b"Some text"
        result = extract_text(data, "NOTES.TXT")
        assert result == "Some text"


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------

class TestDocxExtraction:
    def test_docx_with_minimal_docx(self):
        """Create a minimal valid DOCX in-memory and verify extraction."""
        try:
            from docx import Document
            import io
            doc = Document()
            doc.add_paragraph("Claim 1: A method comprising step A.")
            doc.add_paragraph("Claim 2: The method of claim 1.")
            buf = io.BytesIO()
            doc.save(buf)
            docx_bytes = buf.getvalue()
        except ImportError:
            pytest.skip("python-docx not installed")

        result = extract_text(docx_bytes, "test.docx")
        assert "Claim 1" in result
        assert "Claim 2" in result

    def test_docx_extension_case_insensitive(self):
        try:
            from docx import Document
            import io
            doc = Document()
            doc.add_paragraph("Background of the invention.")
            buf = io.BytesIO()
            doc.save(buf)
            docx_bytes = buf.getvalue()
        except ImportError:
            pytest.skip("python-docx not installed")

        result = extract_text(docx_bytes, "PATENT.DOCX")
        assert "Background" in result


# ---------------------------------------------------------------------------
# Unsupported file type
# ---------------------------------------------------------------------------

class TestUnsupportedType:
    def test_raises_value_error_for_png(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text(b"\x89PNG\r\n", "image.png")

    def test_raises_value_error_for_no_extension(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text(b"data", "mysterious_file")


# ---------------------------------------------------------------------------
# MIN_CHARS_PER_PAGE constant sanity check
# ---------------------------------------------------------------------------

def test_min_chars_per_page_is_positive():
    assert MIN_CHARS_PER_PAGE > 0
